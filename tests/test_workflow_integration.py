from pathlib import Path

from docx import Document

from app.agent.orchestrator import AgentOrchestrator
from app.models.base import BaseModelProvider
from app.models.manager import ModelManager
from app.tools.documents import DocumentTool
from app.artifacts.docx import DOCXArtifactGenerator


class FakeProvider(BaseModelProvider):
    """
    Deterministic provider for integration tests.
    """

    def __init__(self):
        self.calls = []

    def generate(
        self,
        model_name,
        prompt,
        system_prompt=None,
    ):
        self.calls.append(
            {
                "type": "text",
                "model": model_name,
                "prompt": prompt,
            }
        )

        return (
            f"Generated response from {model_name}"
        )

    def generate_multimodal(
        self,
        model_name,
        prompt,
        images,
        system_prompt=None,
    ):
        self.calls.append(
            {
                "type": "multimodal",
                "model": model_name,
                "prompt": prompt,
                "images": images,
            }
        )

        return (
            f"Visual response from {model_name}"
        )

    def is_available(self):
        return True


def test_document_tool_uses_extract(tmp_path: Path):
    """
    Verify that the agent-facing document tool correctly
    calls the existing DocumentPipeline.extract() API.
    """

    document = tmp_path / "report.txt"

    document.write_text(
        "Inspection report: Pump P-101 operating normally.",
        encoding="utf-8",
    )

    tool = DocumentTool()

    result = tool.execute(
        file_path=str(document)
    )

    assert result["status"] == "processed"
    assert result["file_type"] == "txt"
    assert "P-101" in result["text"]


def test_orchestrator_accepts_file_path(tmp_path: Path):
    """
    Verify that a local document path reaches AgentState
    through the public orchestrator interface.
    """

    document = tmp_path / "inspection.txt"

    document.write_text(
        "Inspection finding: P-101 requires review.",
        encoding="utf-8",
    )

    provider = FakeProvider()

    manager = ModelManager(
        provider=provider
    )

    agent = AgentOrchestrator(
        model_manager=manager
    )

    state = agent.run(
        user_input="Analyze this inspection report.",
        file_path=str(document),
    )

    assert state.metadata["file_path"] == str(document)

    assert (
        "documents"
        in state.metadata["selected_tools"]
    )

    assert "documents" in state.tool_results

    assert state.completed is True
    assert state.error is None


def test_docx_artifact_is_valid():
    generator = DOCXArtifactGenerator()

    result = generator.generate_approval_note(
        title="Inspection Approval Note",
        summary="Inspection completed.",
        findings=[
            "Pump P-101 requires review."
        ],
        recommendations=[
            "Inspect pump before continued operation."
        ],
        uncertainties=[
            "Final field verification required."
        ],
        source_document="inspection_report.pdf",
        filename="test_approval_note.docx",
    )

    assert result["status"] == "created"

    output = Path(
        "data/artifacts/test_approval_note.docx"
    )

    assert output.exists()
    assert output.stat().st_size > 0

    document = Document(output)

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
    ]

    combined_text = "\n".join(paragraphs)

    assert "Inspection Approval Note" in combined_text
    assert "Inspection completed." in combined_text
    assert "Pump P-101 requires review." in combined_text

    assert len(document.tables) == 1
def test_artifact_tool_creates_valid_docx():
    from app.tools.artifact import ArtifactTool

    output = Path(
        "data/artifacts/test_validated_note.docx"
    )

    tool = ArtifactTool()

    result = tool.execute(
        artifact_type="docx",
        title="Approval Note",
        summary="Inspection summary.",
        findings=[
            "Finding A"
        ],
        recommendations=[
            "Recommendation A"
        ],
        uncertainties=[
            "Human review required."
        ],
        source_document="inspection.txt",
        filename="test_validated_note.docx",
    )

    assert result["status"] == "created"
    assert result["artifact_type"] == "docx"
    assert result["path"] == str(output)

    assert output.exists()
    assert output.stat().st_size > 0

    assert result["validation"]["valid"] is True